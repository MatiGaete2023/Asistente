#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
motor/procesador.py — v8.6
Optimizaciones:
  CLAIM-3: _guardar_excel usa dataframe_to_rows (3x más rápido)
  CLAIM-4: _construir_indice_hoja2 usa to_dict('records') (8.6x más rápido)
"""

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows   # CLAIM-3

from .utilidades import (
    normalizar_match, detectar_tribunal, titulo_programa,
    prefijo_observacion, fecha_es, get_date, audiencia_suffix
)
from .mapeo_columnas import mapear_columnas, mapear_columnas_hoja2
from .textos import render
from .reglas_espera import generar_observacion_espera
from .reglas_cumplimiento import (
    generar_observacion_cumplimiento,
    r5_aplica_para_fila,
    medida_vencida_para_fila
)
from .reglas_informes import generar_observacion_informes

logger = logging.getLogger(__name__)


# ── cruce Hoja2 ───────────────────────────────────────────────────────────────

def _construir_indice_hoja2(df_h2, cols_h2):
    """
    CLAIM-4: to_dict('records') en vez de iterrows — 8.6x más rápido.
    """
    indice = {}
    campos = ["rit", "rut", "nombre", "tribunal", "programa", "vencimiento"]
    cols_v = [cols_h2.get(c) for c in campos]
    if any(c is None for c in cols_v):
        logger.warning("Hoja2: faltan columnas — cruce desactivado")
        return {}
    col_rit, col_rut, col_nom, col_tri, col_prog, col_venc = cols_v
    hoy = datetime.now().date()

    for row in df_h2.to_dict("records"):          # CLAIM-4
        venc = get_date(row.get(col_venc))
        if venc is None:
            continue
        vd = venc.date() if hasattr(venc, "date") else venc
        if vd <= hoy:  # descartar vencidas incluyendo hoy
            continue
        rit = normalizar_match(row.get(col_rit))
        if not rit:
            continue
        indice[(
            rit,
            normalizar_match(row.get(col_rut)),
            normalizar_match(row.get(col_nom)),
            normalizar_match(row.get(col_tri)),
            normalizar_match(row.get(col_prog)),
        )] = vd
    return indice


def _clave_fila_h1(row, cols):
    keys = ("rit", "rut", "nombre", "tribunal", "programa")
    if not all(cols.get(k) for k in keys):
        return None
    return tuple(normalizar_match(row.get(cols[k], "")) for k in keys)


def _obs_cruce(row, cols, fecha_venc):
    programa = str(row.get(cols.get("programa"), "")).strip()
    nombre   = str(row.get(cols.get("nombre"),   "")).strip()
    pfx      = prefijo_observacion(nombre, programa)
    prog_fmt = titulo_programa(programa)
    aud      = audiencia_suffix(row, cols)
    obs = render("CUMPLIMIENTO", "CRUCE_HOJA2", PROGRAMA=prog_fmt, FECHA_VENCIMIENTO=fecha_es(fecha_venc))
    base = obs[:-1] if (aud and obs.endswith(".")) else obs
    return pfx + base + aud


# ── guardado Excel ─────────────────────────────────────────────────────────────

def _guardar_excel(df, modo, ruta, q):
    """
    CLAIM-3: dataframe_to_rows en vez de iterrows — 3x más rápido.
    """
    try:
        Path(ruta).mkdir(parents=True, exist_ok=True)
        nombre = f"RUS_{modo}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        arch   = Path(ruta) / nombre
        wb = Workbook()
        ws = wb.active
        ws.title = modo

        # CLAIM-3: escritura en bloque
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

        # Estilos del encabezado (fila 1)
        for c in ws[1]:
            c.font      = Font(bold=True, color="FFFFFF")
            c.fill      = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")

        # Ancho de columnas
        for col in ws.columns:
            ancho = max((len(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(ancho + 2, 60)

        wb.save(arch)
        q.put(("log", f"✅ Guardado: {nombre}"))
        q.put(("log", f"📁 {arch}"))
        return nombre
    except Exception as e:
        q.put(("log", f"❌ Error al guardar: {e}"))
        return None


# ── procesadores ──────────────────────────────────────────────────────────────

def _calcular_simple(df, cols, fn_regla):
    """
    Calcula el DataFrame con columna OBSERVACION para ESPERA/INFORMES,
    SIN guardar a disco. Reutilizado por _procesar_simple (que sí guarda)
    y por calcular_preview() (vista previa GUI — S3, no escribe archivo).
    """
    col_trib = cols.get("tribunal")
    if "OBSERVACION" in df.columns:
        df = df.drop(columns=["OBSERVACION"])

    def aplicar(row):
        try:
            trib = detectar_tribunal(str(row.get(col_trib, ""))) or "LAJA"
            return fn_regla(row, trib, cols)
        except Exception as e:
            return f"ERROR: {e}"

    df_r = df.copy()
    df_r["OBSERVACION"] = df_r.apply(aplicar, axis=1)
    return df_r


def _procesar_simple(df, modo, fn_regla, config, q):
    """Procesador genérico para ESPERA e INFORMES."""
    cols     = mapear_columnas(df, modo)
    col_prog = cols.get("programa")
    col_trib = cols.get("tribunal")
    if not col_prog or not col_trib:
        q.put(("log", "❌ Columnas DERIVACION o TRIBUNAL no encontradas"))
        q.put(("done", f"❌ {modo} cancelado: columnas no encontradas."))
        return

    df_r = _calcular_simple(df, cols, fn_regla)
    nombre = _guardar_excel(df_r, modo, config["ruta_salida_excel"], q)
    if nombre:
        q.put(("done", f"{modo} completado.\n{len(df_r)} casos.\nArchivo: {nombre}"))
    else:
        q.put(("done", f"❌ {modo}: error al guardar el archivo."))


def _calcular_cumplimiento(df_h1, cols, indice):
    """
    Calcula el DataFrame con columna OBSERVACION para CUMPLIMIENTO,
    incluyendo el cruce con Hoja2 (indice), SIN guardar a disco.
    Reutilizado por _procesar_cumplimiento (que sí guarda) y por
    calcular_preview() (vista previa GUI — S3, no escribe archivo).
    """
    col_trib = cols.get("tribunal")
    if "OBSERVACION" in df_h1.columns:
        df_h1 = df_h1.drop(columns=["OBSERVACION"])

    def aplicar(row):
        try:
            trib = detectar_tribunal(str(row.get(col_trib, ""))) or "LAJA"
            obs  = generar_observacion_cumplimiento(row, trib, cols)
            if (indice
                    and not medida_vencida_para_fila(row, cols)
                    and not r5_aplica_para_fila(row, cols)):
                clave = _clave_fila_h1(row, cols)
                if clave and clave in indice:
                    obs = _obs_cruce(row, cols, indice[clave])
            return obs
        except Exception as e:
            return f"ERROR: {e}"

    df_r = df_h1.copy()
    df_r["OBSERVACION"] = df_r.apply(aplicar, axis=1)
    return df_r


def _procesar_cumplimiento(df_h1, df_h2, config, q):
    cols     = mapear_columnas(df_h1, "CUMPLIMIENTO")
    col_prog = cols.get("programa")
    col_trib = cols.get("tribunal")
    if not col_prog or not col_trib:
        q.put(("log", "❌ Columnas DERIVACION o TRIBUNAL no encontradas"))
        q.put(("done", "❌ CUMPLIMIENTO cancelado: columnas no encontradas."))
        return

    indice = {}
    if df_h2 is not None:
        indice = _construir_indice_hoja2(df_h2, mapear_columnas_hoja2(df_h2))
        q.put(("log", f"📋 Hoja2: {len(indice)} registros con fecha futura"))
    else:
        q.put(("log", "ℹ️  Sin Hoja2 — cruce desactivado"))

    df_r = _calcular_cumplimiento(df_h1, cols, indice)
    nombre = _guardar_excel(df_r, "CUMPLIMIENTO", config["ruta_salida_excel"], q)
    if nombre:
        q.put(("done", f"CUMPLIMIENTO completado.\n{len(df_r)} casos.\nArchivo: {nombre}"))
    else:
        q.put(("done", "❌ CUMPLIMIENTO: error al guardar el archivo."))


# ── dispatcher ────────────────────────────────────────────────────────────────

def _leer_excel(path):
    engine = "xlrd" if str(path).endswith(".xls") else "openpyxl"
    df = pd.read_excel(path, engine=engine)
    df.columns = [str(x).strip() for x in df.columns]
    return df.dropna(how="all").fillna("")


def procesar(df_o_path, modo, config, q):
    if modo == "CUMPLIMIENTO":
        if isinstance(df_o_path, (str, Path)):
            path = str(df_o_path)
            try:
                engine   = "xlrd" if path.endswith(".xls") else "openpyxl"
                xl       = pd.ExcelFile(path, engine=engine)
                names    = {s.strip().lower(): s for s in xl.sheet_names}
                hoja_key = "cumplimiento" if "cumplimiento" in names else list(names.keys())[0]
                df_h1    = pd.read_excel(xl, sheet_name=names[hoja_key])
                df_h1.columns = [str(x).strip() for x in df_h1.columns]
                df_h1    = df_h1.dropna(how="all").fillna("")

                otras = [s for s in xl.sheet_names if s.strip().lower() != hoja_key]
                df_h2 = None
                if otras:
                    try:
                        df_h2 = pd.read_excel(xl, sheet_name=otras[0])
                        df_h2.columns = [str(x).strip() for x in df_h2.columns]
                        df_h2 = df_h2.dropna(how="all").fillna("")
                        q.put(("log", f"📄 Hoja2: '{otras[0]}' ({len(df_h2)} filas)"))
                    except Exception as e:
                        q.put(("log", f"⚠️  No se pudo cargar Hoja2: {e}"))

                q.put(("log", f"✓ {len(df_h1)} filas CUMPLIMIENTO"))
                _procesar_cumplimiento(df_h1, df_h2, config, q)
            except Exception as e:
                q.put(("log", f"❌ Error cargando archivo: {e}"))
                q.put(("done", f"❌ CUMPLIMIENTO cancelado: {e}"))
        else:
            df = df_o_path
            df.columns = [str(x).strip() for x in df.columns]
            _procesar_cumplimiento(df.dropna(how="all").fillna(""), None, config, q)

    elif modo in ("ESPERA", "INFORMES"):
        fn = generar_observacion_espera if modo == "ESPERA" else generar_observacion_informes
        try:
            df = (_leer_excel(df_o_path)
                  if isinstance(df_o_path, (str, Path))
                  else df_o_path.dropna(how="all").fillna(""))
            q.put(("log", f"✓ {len(df)} filas {modo}"))
            _procesar_simple(df, modo, fn, config, q)
        except Exception as e:
            q.put(("log", f"❌ Error cargando archivo: {e}"))
            q.put(("done", f"❌ {modo} cancelado: {e}"))
    else:
        q.put(("log", f"❌ Modo desconocido: {modo}"))
        q.put(("done", f"❌ Modo desconocido: {modo}"))


# ── vista previa (S3 — v8.14) ──────────────────────────────────────────────────

def calcular_preview(df_o_path, modo):
    """
    Calcula el DataFrame con columna OBSERVACION para `modo`, IDENTICO
    byte-a-byte al que produce procesar() (misma logica, mismos textos,
    mismo cruce Hoja2 en CUMPLIMIENTO), pero SIN escribir ningun archivo
    a disco. Usado exclusivamente por la vista previa de gui/app.py (S3).

    Retorna (df_resultado, None) en exito, o (None, mensaje_error) si
    faltan columnas obligatorias o el modo es desconocido.
    """
    try:
        if modo == "CUMPLIMIENTO":
            if isinstance(df_o_path, (str, Path)):
                path   = str(df_o_path)
                engine = "xlrd" if path.endswith(".xls") else "openpyxl"
                xl     = pd.ExcelFile(path, engine=engine)
                names  = {s.strip().lower(): s for s in xl.sheet_names}
                hoja_key = "cumplimiento" if "cumplimiento" in names else list(names.keys())[0]
                df_h1  = pd.read_excel(xl, sheet_name=names[hoja_key])
                df_h1.columns = [str(x).strip() for x in df_h1.columns]
                df_h1  = df_h1.dropna(how="all").fillna("")

                otras = [s for s in xl.sheet_names if s.strip().lower() != hoja_key]
                df_h2 = None
                if otras:
                    try:
                        df_h2 = pd.read_excel(xl, sheet_name=otras[0])
                        df_h2.columns = [str(x).strip() for x in df_h2.columns]
                        df_h2 = df_h2.dropna(how="all").fillna("")
                    except Exception:
                        df_h2 = None
            else:
                df_h1 = df_o_path.copy()
                df_h1.columns = [str(x).strip() for x in df_h1.columns]
                df_h1 = df_h1.dropna(how="all").fillna("")
                df_h2 = None

            cols = mapear_columnas(df_h1, "CUMPLIMIENTO")
            if not cols.get("programa") or not cols.get("tribunal"):
                return None, "Columnas DERIVACION o TRIBUNAL no encontradas"

            indice = {}
            if df_h2 is not None:
                indice = _construir_indice_hoja2(df_h2, mapear_columnas_hoja2(df_h2))

            df_r = _calcular_cumplimiento(df_h1, cols, indice)
            return df_r, None

        elif modo in ("ESPERA", "INFORMES"):
            fn = generar_observacion_espera if modo == "ESPERA" else generar_observacion_informes
            df = (_leer_excel(df_o_path)
                  if isinstance(df_o_path, (str, Path))
                  else df_o_path.dropna(how="all").fillna(""))
            cols = mapear_columnas(df, modo)
            if not cols.get("programa") or not cols.get("tribunal"):
                return None, "Columnas DERIVACION o TRIBUNAL no encontradas"

            df_r = _calcular_simple(df, cols, fn)
            return df_r, None

        else:
            return None, f"Modo desconocido: {modo}"

    except Exception as e:
        return None, f"Error al calcular vista previa: {e}"
