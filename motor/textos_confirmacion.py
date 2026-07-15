#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
motor/textos_confirmacion.py — v9.0.1

Flujo de confirmación de textos que NO depende de sesiones de chat con
una IA. Resuelve F-7 estructuralmente:

  1. exportar_pendientes(ruta_xlsx) -> escribe un .xlsx con los textos
     que tienen confirmado=False actualmente en
     motor/textos_observaciones.json.

  2. La persona usuaria abre ese Excel y, por fila:
       - Si el texto está bien tal cual: deja TEXTO_CORRECTO vacío y
         escribe SI en APROBADO.
       - Si el texto necesita cambios: escribe la redacción exacta en
         TEXTO_CORRECTO (y opcionalmente SI en APROBADO si ya la da
         por buena así).

  3. importar_confirmaciones(ruta_xlsx) -> aplica esos cambios a
     motor/textos_observaciones.json:
       - TEXTO_CORRECTO no vacío -> reemplaza "texto" (COPY-PASTE
         literal de lo que escribió la persona usuaria, NUNCA parafraseado).
       - APROBADO == "SI" -> marca "confirmado": true.
     Hace backup con timestamp del JSON antes de sobrescribirlo.

Restricción inviolable del proyecto (ver memoria / PLAN v8.14 §2.1):
el texto que termina en el JSON es SIEMPRE lo que el usuario escribió,
verbatim. Este módulo no parafrasea, no correige ortografía, no agrega
ni quita puntuación.
"""

import json
import os
import shutil
import string
import tempfile
from collections import Counter
from threading import RLock
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from .textos import _RUTA
from .utilidades import validar_archivo_excel

COLUMNAS = ["MODO", "ID", "TEXTO_ACTUAL", "TEXTO_CORRECTO", "APROBADO"]
_IMPORT_LOCK = RLock()


def _cargar_json(ruta_json: Path) -> dict:
    with open(ruta_json, encoding="utf-8") as f:
        return json.load(f)


def _listar_no_confirmados_de(data: dict):
    out = []
    for modo, reglas in data.items():
        if modo == "_meta":
            continue
        for id_regla, entry in reglas.items():
            if not entry.get("confirmado", False):
                out.append((modo, id_regla))
    return out


def contar_pendientes(ruta_json=None) -> int:
    """Cantidad de textos con confirmado=False. Usado para el badge del
    botón en la GUI: '📋 Textos pendientes (N)'."""
    ruta_json = Path(ruta_json) if ruta_json else _RUTA
    data = _cargar_json(ruta_json)
    return len(_listar_no_confirmados_de(data))


def exportar_pendientes(ruta_xlsx, ruta_json=None) -> int:
    """Escribe TEXTOS_PENDIENTES.xlsx. Retorna cantidad de filas."""
    ruta_json = Path(ruta_json) if ruta_json else _RUTA
    ruta_xlsx = Path(ruta_xlsx)
    if ruta_xlsx.suffix.lower() != ".xlsx":
        raise ValueError("La exportación de textos debe usar formato .xlsx")
    data = _cargar_json(ruta_json)
    pendientes = _listar_no_confirmados_de(data)

    filas = []
    for modo, id_regla in pendientes:
        entry = data[modo][id_regla]
        filas.append({
            "MODO": modo,
            "ID": id_regla,
            "TEXTO_ACTUAL": entry["texto"],
            "TEXTO_CORRECTO": "",
            "APROBADO": "",
        })

    # openpyxl interpreta automáticamente cadenas que comienzan con '=' como
    # fórmulas. Se fuerza el tipo string para conservar el texto literal y
    # permitir que la planilla vuelva a importarse sin alterar TEXTO_ACTUAL.
    ruta_xlsx.parent.mkdir(parents=True, exist_ok=True)
    fd, temporal_nombre = tempfile.mkstemp(
        prefix=f".{ruta_xlsx.stem}.", suffix=".tmp.xlsx", dir=ruta_xlsx.parent
    )
    os.close(fd)
    temporal = Path(temporal_nombre)
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "TEXTOS"
        for columna, nombre in enumerate(COLUMNAS, start=1):
            celda = ws.cell(row=1, column=columna, value=nombre)
            celda.data_type = "s"
        for fila_excel, fila in enumerate(filas, start=2):
            for columna, nombre in enumerate(COLUMNAS, start=1):
                valor = fila[nombre]
                celda = ws.cell(row=fila_excel, column=columna, value=valor)
                if isinstance(valor, str):
                    celda.data_type = "s"
        wb.save(temporal)
        with open(temporal, "rb") as f:
            os.fsync(f.fileno())
        os.replace(temporal, ruta_xlsx)
    except Exception:
        temporal.unlink(missing_ok=True)
        raise
    return len(filas)


def _backup_json(ruta_json: Path) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    destino = ruta_json.parent / f"{ruta_json.stem}.backup_{ts}{ruta_json.suffix}"
    shutil.copy2(ruta_json, destino)
    return destino


def _celda_texto(valor, *, verbatim=False) -> str:
    """Convierte una celda sin transformar NaN en la cadena literal 'nan'."""
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    texto = str(valor)
    return texto if verbatim else texto.strip()


def _placeholders(texto: str) -> Counter:
    """Extrae la firma completa de formato y rechaza accesos indirectos."""
    try:
        firma = Counter()
        for _, campo, formato, conversion in string.Formatter().parse(texto):
            if campo is None:
                continue
            if not campo or not campo.replace("_", "").isalnum():
                raise ValueError(f"placeholder no permitido: {campo!r}")
            firma[(campo, formato, conversion)] += 1
        return firma
    except ValueError as exc:
        raise ValueError(f"llaves de formato inválidas: {exc}") from exc


def _validar_reemplazo(texto_actual: str, texto_nuevo: str) -> str | None:
    if not texto_nuevo.strip():
        return "el texto corregido está vacío"
    try:
        actuales = _placeholders(texto_actual)
        nuevos = _placeholders(texto_nuevo)
    except ValueError as exc:
        return str(exc)
    if actuales != nuevos:
        return ("los placeholders deben conservarse exactamente; "
                f"esperados={list(actuales.elements())}, "
                f"recibidos={list(nuevos.elements())}")
    return None


def _guardar_json_atomico(data: dict, ruta_json: Path) -> None:
    """Escribe y sincroniza un temporal antes de reemplazar el catálogo."""
    ruta_json.parent.mkdir(parents=True, exist_ok=True)
    fd, temporal = tempfile.mkstemp(
        prefix=f".{ruta_json.stem}.", suffix=".tmp", dir=ruta_json.parent
    )
    temporal_path = Path(temporal)
    try:
        if ruta_json.exists():
            os.chmod(temporal_path, ruta_json.stat().st_mode & 0o777)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.write("\n")
            f.flush()
            os.fsync(f.fileno())
        os.replace(temporal_path, ruta_json)
    except Exception:
        try:
            os.close(fd)
        except OSError:
            pass
        temporal_path.unlink(missing_ok=True)
        raise


def importar_confirmaciones(ruta_xlsx, ruta_json=None) -> dict:
    """
    Aplica TEXTOS_PENDIENTES.xlsx (ya editado) a
    motor/textos_observaciones.json. Retorna:
        {
          "actualizados": [ "MODO.ID", ... ],   # texto reemplazado
          "confirmados":  [ "MODO.ID", ... ],   # confirmado=true
          "sin_cambios":  [ "MODO.ID", ... ],
          "errores":      [ "mensaje", ... ],
          "backup":       "ruta al JSON previo respaldado",
        }
    Lanza ValueError si al Excel le faltan columnas requeridas.
    """
    ruta_json = Path(ruta_json) if ruta_json else _RUTA
    with _IMPORT_LOCK:
        df = pd.read_excel(validar_archivo_excel(ruta_xlsx))
        faltan = [c for c in COLUMNAS if c not in df.columns]
        if faltan:
            raise ValueError(f"Faltan columnas en el Excel: {faltan}")

        data = _cargar_json(ruta_json)
        resumen = {"actualizados": [], "confirmados": [], "sin_cambios": [],
                   "errores": [], "backup": None}
        hubo_cambios = False

        for _, row in df.iterrows():
            modo = _celda_texto(row.get("MODO", ""))
            id_regla = _celda_texto(row.get("ID", ""))
            texto_excel = _celda_texto(row.get("TEXTO_ACTUAL", ""), verbatim=True)
            texto_correcto = _celda_texto(row.get("TEXTO_CORRECTO", ""), verbatim=True)
            aprobado = _celda_texto(row.get("APROBADO", "")).upper()
            clave = f"{modo}.{id_regla}"

            if not modo or not id_regla:
                continue
            if modo not in data or id_regla not in data.get(modo, {}):
                resumen["errores"].append(f"{clave}: no existe en el JSON actual")
                continue

            entry = data[modo][id_regla]
            if entry.get("confirmado", False):
                resumen["errores"].append(
                    f"{clave}: la regla ya estaba confirmada; exporta una planilla vigente"
                )
                continue
            if texto_excel != entry.get("texto", ""):
                resumen["errores"].append(
                    f"{clave}: TEXTO_ACTUAL no coincide con el catálogo vigente"
                )
                continue
            if aprobado not in ("", "SI", "NO"):
                resumen["errores"].append(
                    f"{clave}: APROBADO debe ser SI, NO o quedar vacío"
                )
                continue

            cambiado = False
            if texto_correcto.strip():
                error = _validar_reemplazo(entry.get("texto", ""), texto_correcto)
                if error:
                    resumen["errores"].append(f"{clave}: {error}")
                    continue
                entry["texto"] = texto_correcto  # literal; no strip ni paráfrasis
                resumen["actualizados"].append(clave)
                cambiado = True

            if aprobado == "SI":
                entry["confirmado"] = True
                resumen["confirmados"].append(clave)
                cambiado = True

            if cambiado:
                hubo_cambios = True
            else:
                resumen["sin_cambios"].append(clave)

        if hubo_cambios:
            backup = _backup_json(ruta_json)
            _guardar_json_atomico(data, ruta_json)
            resumen["backup"] = str(backup)

            if ruta_json == _RUTA:
                # Reflejar el cambio en la misma sesión sin una carrera sobre
                # la referencia global del catálogo.
                import motor.textos as _textos_mod
                _textos_mod.invalidar_cache()

        return resumen
