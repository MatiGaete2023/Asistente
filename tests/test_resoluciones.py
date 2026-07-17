import pandas as pd
from docx import Document

from resoluciones.generador_resoluciones import fecha_numerica, generar_resoluciones


OBS_PC_IE = (
    "Se genera proyecto de resolución pidiendo cuenta al programa "
    "respecto del ingreso efectivo"
)


def test_resolucion_valida_no_contiene_marcadores(tmp_path):
    entrada = tmp_path / "entrada.xlsx"
    pd.DataFrame([{
        "OBSERVACION": OBS_PC_IE,
        "TRIBUNAL": "MULCHEN",
        "RIT": "R-1",
        "NOMBRE": "CASO SINTÉTICO",
        "RUT": "1-9",
        "DERIVACION": "PRM SINTÉTICO",
        "DURACION": "6 mes(es)",
        "FEC. RESOLUCIÓN": "01/07/2026",
    }]).to_excel(entrada, index=False)

    resultado = generar_resoluciones(str(entrada), str(tmp_path))
    assert resultado["total_resoluciones"] == 1
    documento = Document(resultado["archivo_generado"])
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)
    assert "COMPLETAR" not in texto
    assert "[ERROR" not in texto
    assert "[SIN PLANTILLA" not in texto


def test_resolucion_incompleta_va_al_word_con_marcadores_completar(tmp_path):
    """Flujo aprobado (Mejoras §2.1: mecanismo intacto): los datos ausentes
    van como COMPLETAR para edición manual — la fila nunca se omite."""
    entrada = tmp_path / "incompleta.xlsx"
    pd.DataFrame([{
        "OBSERVACION": OBS_PC_IE,
        "TRIBUNAL": "MULCHEN",
        "RIT": "R-2",
        "NOMBRE": "Caso",
    }]).to_excel(entrada, index=False)

    resultado = generar_resoluciones(str(entrada), str(tmp_path))
    assert resultado["total_resoluciones"] == 1
    assert resultado["archivo_generado"] is not None
    documento = Document(resultado["archivo_generado"])
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)
    assert "COMPLETAR" in texto


def test_tribunal_sin_plantilla_genera_bloque_manual(tmp_path):
    """Tomé no tiene plantilla Word: entra al documento como bloque
    [SIN PLANTILLA] y queda listado en faltantes (comportamiento aprobado)."""
    entrada = tmp_path / "tome.xlsx"
    pd.DataFrame([{
        "OBSERVACION": OBS_PC_IE,
        "TRIBUNAL": "TOME",
        "RIT": "R-3",
        "NOMBRE": "Caso",
        "RUT": "1-9",
        "DERIVACION": "PIE",
    }]).to_excel(entrada, index=False)

    resultado = generar_resoluciones(str(entrada), str(tmp_path))
    assert resultado["archivo_generado"] is not None
    assert resultado["faltantes"] and resultado["faltantes"][0]["tipo"] == "PC_IE"
    documento = Document(resultado["archivo_generado"])
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)
    assert "[SIN PLANTILLA" in texto
    assert "R-3" in texto


def test_fecha_invalida_produce_marcador_completar():
    assert fecha_numerica("no-es-fecha") == "COMPLETAR"


def test_resoluciones_desde_archivo_origen_prioriza_pedir_cuenta(tmp_path):
    """Un lote mixto prioriza T ESPERA para pedir cuenta sobre NOMENCL."""
    entrada = tmp_path / "origen.xlsx"
    pd.DataFrame([{
        "OBSERVACION": "Aplicar nomenclaturas a fin de regularizar",
        "TRIBUNAL": "MULCHEN", "RIT": "R-ORIGEN", "NOMBRE": "Caso",
        "RUT": "1-9", "DERIVACION": "PRM", "DURACION": "6 mes(es)",
        "FEC. RESOLUCIÓN": "01/07/2026", "T ESPERA": 40,
    }]).to_excel(entrada, index=False)

    resultado = generar_resoluciones(str(entrada), str(tmp_path))
    assert resultado["total_resoluciones"] == 1
    texto = "\n".join(p.text for p in Document(resultado["archivo_generado"]).paragraphs)
    assert "pídase cuenta al programa" in texto
