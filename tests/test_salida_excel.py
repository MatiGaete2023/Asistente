# -*- coding: utf-8 -*-
"""Columnas de salida §8 del catálogo, hoja VALIDACION (D6) y G-06 sin default."""
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from motor.mapeo_columnas import mapear_columnas
from motor.procesador import _calcular_simple, _guardar_excel
from motor.reglas_espera import generar_observacion_espera


class _Cola:
    def __init__(self):
        self.mensajes = []

    def put(self, m):
        self.mensajes.append(m)


def _df_espera(tribunal="LAJA", espera=40):
    return pd.DataFrame([{
        "RIT": "P-1-2026", "RUT": "1-9", "NOMBRE": "ANA SOTO",
        "TRIBUNAL": tribunal, "DERIVACION": "PIE Sintético",
        "T ESPERA": espera,
    }])


def _calcular(df):
    cols = mapear_columnas(df, "ESPERA")
    return _calcular_simple(df, cols, generar_observacion_espera)


def test_orden_columnas_fecha_obs_tt_cc_res():
    df_r, _ = _calcular(_df_espera())
    columnas = list(df_r.columns)
    i = columnas.index("OBSERVACION")
    assert columnas[i - 1] == "FECHA_OBS"
    assert columnas[i + 1:i + 4] == ["TT", "CC", "RES"]


def test_fecha_obs_automatica_y_tt_cc_res_vacias():
    df_r, _ = _calcular(_df_espera())
    assert df_r["FECHA_OBS"].iloc[0] == datetime.now().strftime("%d/%m/%Y")
    for col in ("TT", "CC", "RES"):
        assert df_r[col].iloc[0] == ""  # vacías, no 0 (D2 / catálogo §13)


def test_reproceso_no_duplica_columnas():
    df_r, _ = _calcular(_df_espera())
    df_r2, _ = _calcular(df_r)
    assert list(df_r2.columns).count("FECHA_OBS") == 1
    assert list(df_r2.columns).count("OBSERVACION") == 1


def test_g06_tribunal_desconocido_sin_default_laja():
    """Antes v9 heredaba 'LAJA' (H-01): 40 días → proyecto de resolución.
    Ahora: E-06 sin proyecto + incidencia G-06."""
    df_r, incidencias = _calcular(_df_espera(tribunal="JUZGADO DE COYHAIQUE"))
    o = df_r["OBSERVACION"].iloc[0]
    assert "proyecto de resolución" not in o
    assert "a la espera de ingreso efectivo al programa" in o
    regs = incidencias.como_dataframe().to_dict("records")
    assert any(r["REGLA"] == "G-06" for r in regs)


def test_guardar_escribe_hoja_validacion(tmp_path):
    df_r, incidencias = _calcular(_df_espera(tribunal="OTRO TRIBUNAL"))
    q = _Cola()
    nombre = _guardar_excel(df_r, "ESPERA", str(tmp_path), q, incidencias)
    wb = load_workbook(tmp_path / nombre)
    assert "VALIDACION" in wb.sheetnames
    filas = list(wb["VALIDACION"].iter_rows(values_only=True))
    assert filas[0] == ("FILA_EXCEL", "RIT", "REGLA", "MOTIVO")
    assert len(filas) >= 2


def test_guardar_sin_incidencias_no_crea_hoja(tmp_path):
    df_r, incidencias = _calcular(_df_espera())
    q = _Cola()
    nombre = _guardar_excel(df_r, "ESPERA", str(tmp_path), q, incidencias)
    wb = load_workbook(tmp_path / nombre)
    assert "VALIDACION" not in wb.sheetnames
