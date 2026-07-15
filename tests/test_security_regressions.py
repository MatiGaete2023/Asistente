import json
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from comunicaciones.contactos_programas import CatastroContactos
from comunicaciones.generador_correos import (
    CC_FIJO,
    GeneradorCorreos,
    PATRON_ESPERA,
    _fila_fecha,
)
from motor.composicion import Incidencias
from motor.mapeo_columnas import mapear_columnas
from motor.procesador import _construir_indice_hoja2, _guardar_excel, calcular_preview
from motor.reglas_informes import generar_observacion_informes
from motor.utilidades import get_int, validar_archivo_excel


CATASTRO = Path(__file__).resolve().parents[1] / "comunicaciones" / "catastro_programas.json"


def _programa(catastro, *, dce=False):
    return next(
        registro["nombre"]
        for registro in catastro.todos()
        if ("DCE" in registro["nombre"].upper()) is dce
    )


def _procesar_informes(programa, tribunales):
    borradores = []
    generador = GeneradorCorreos(
        {"cc_fijo": "atacante@example.invalid"},
        str(CATASTRO),
        despachador=lambda borrador: borradores.append(borrador),
    )
    fecha = datetime.now().date() + timedelta(days=5)
    filas = [
        {
            "RIT": f"R-{indice}",
            "TRIBUNAL": tribunal,
            "RUT": "1-9",
            "NOMBRE": "Caso Sintético",
            "DERIVACION": programa,
            "FECHA VENCIMIENTO": fecha,
        }
        for indice, tribunal in enumerate(tribunales, start=1)
    ]
    return generador, generador.procesar(pd.DataFrame(filas)), borradores


def test_mapeo_cumplimiento_incluye_ficha_individual():
    df = pd.DataFrame(columns=["DERIVACION", "TRIBUNAL", "FEC.ACT.F.INDIVIDUAL"])
    assert mapear_columnas(df, "CUMPLIMIENTO")["ficha_ind"] == "FEC.ACT.F.INDIVIDUAL"


def test_incidencias_se_propagan_al_colector_externo():
    incidencias = Incidencias()
    cols = {
        "programa": "DERIVACION",
        "nombre": "NOMBRE",
        "vencimiento": "FECHA VENCIMIENTO",
        "rit": "RIT",
    }
    row = pd.Series({
        "DERIVACION": "PIE Sintético",
        "NOMBRE": "Caso",
        "FECHA VENCIMIENTO": "",
        "RIT": "R-1",
    })
    assert generar_observacion_informes(
        row, "LAJA", cols, incidencias=incidencias, fila_excel=2
    ) == ""
    assert len(incidencias) == 1


def test_contactos_rechazan_consultas_vacias_genericas_y_fuzzy_automatico():
    catastro = CatastroContactos(str(CATASTRO))
    exacto = _programa(catastro)
    assert catastro.resolver(exacto) is not None
    assert catastro.resolver("") is None
    assert catastro.resolver("PRM") is None
    assert catastro.resolver(f"{exacto} texto adicional") is None


def test_variantes_del_mismo_tribunal_generan_un_borrador():
    catastro = CatastroContactos(str(CATASTRO))
    generador, resultado, borradores = _procesar_informes(
        _programa(catastro),
        ["MULCHEN", "Juzgado de Letras y Garantía de Mulchén"],
    )
    assert resultado["borradores_creados"] == 1
    assert len(borradores) == 1
    assert borradores[0]["cc"] == [CC_FIJO]
    assert generador.cc_fijo == CC_FIJO


def test_alias_y_nombre_canonico_se_agrupan_en_un_borrador():
    aliases = json.loads(
        (CATASTRO.parent / "aliases_programas.json").read_text(encoding="utf-8")
    )
    alias = aliases[0]
    catastro = CatastroContactos(str(CATASTRO))
    fecha = datetime.now().date() + timedelta(days=5)
    filas = pd.DataFrame([
        {"RIT": "R-1", "TRIBUNAL": "LAJA", "NOMBRE": "Caso 1",
         "DERIVACION": alias["alias"], "FECHA VENCIMIENTO": fecha},
        {"RIT": "R-2", "TRIBUNAL": "LAJA", "NOMBRE": "Caso 2",
         "DERIVACION": alias["nombre_catastro"], "FECHA VENCIMIENTO": fecha},
    ])
    capturados = []
    generador = GeneradorCorreos(
        {}, str(CATASTRO), despachador=lambda borrador: capturados.append(borrador)
    )
    resultado = generador.procesar(filas)
    assert catastro.resolver(alias["alias"]) is not None
    assert resultado["borradores_creados"] == 1
    assert len(capturados) == 1
    assert capturados[0]["n_registros"] == 2


def test_dce_usa_diagnosticos_en_asunto_y_cuerpo():
    catastro = CatastroContactos(str(CATASTRO))
    _, resultado, borradores = _procesar_informes(_programa(catastro, dce=True), ["LAJA"])
    assert resultado["borradores_creados"] == 1
    assert "diagnósticos" in borradores[0]["asunto"].lower()
    assert "diagnósticos" in borradores[0]["cuerpo_html"].lower()


def test_html_de_excel_se_escapa():
    fila = _fila_fecha(
        '<img src="https://example.invalid/x">',
        "LAJA",
        "1-9",
        "<b>Caso</b>",
        datetime.now(),
    )
    assert "<img" not in fila
    assert "<b>" not in fila
    assert "&lt;img" in fila
    assert "&lt;b&gt;" in fila


def test_espera_sin_observacion_no_genera_borradores():
    catastro = CatastroContactos(str(CATASTRO))
    programa = _programa(catastro)
    generador = GeneradorCorreos({}, str(CATASTRO), despachador=lambda _: None)
    resultado = generador.procesar_espera(pd.DataFrame([{
        "RIT": "R-1", "TRIBUNAL": "LAJA", "NOMBRE": "Caso",
        "DERIVACION": programa, "T ESPERA": 40,
    }]))
    assert resultado["borradores_creados"] == 0
    assert any("OBSERVACION" in error for error in resultado["errores"])


def test_correos_degradan_suavemente_ante_filas_invalidas():
    """D9: una fila defectuosa se excluye con aviso pero el resto del lote
    SIGUE generando borradores — nunca se aborta el lote completo."""
    catastro = CatastroContactos(str(CATASTRO))
    programa = _programa(catastro)
    capturados = []
    generador = GeneradorCorreos(
        {}, str(CATASTRO), despachador=lambda b: capturados.append(b)
    )
    buena = {
        "RIT": "R-1", "TRIBUNAL": "LAJA", "NOMBRE": "Caso",
        "DERIVACION": programa,
        "FECHA VENCIMIENTO": datetime.now() + timedelta(days=5),
    }
    mala = dict(buena, RIT="R-2", **{"FECHA VENCIMIENTO": "no-es-fecha"})
    resultado = generador.procesar(pd.DataFrame([buena, mala]))
    assert resultado["borradores_creados"] == 1
    assert any("FECHA VENCIMIENTO" in error for error in resultado["errores"])

    # Tribunal no reconocido: se agrupa por el valor crudo con advertencia,
    # sin perder la fila.
    rara = dict(buena, RIT="R-3", TRIBUNAL="DESCONOCIDO")
    resultado = generador.procesar(pd.DataFrame([rara]))
    assert resultado["borradores_creados"] == 1
    assert any("TRIBUNAL" in error for error in resultado["errores"])


def test_espera_rechaza_dias_invalidos():
    catastro = CatastroContactos(str(CATASTRO))
    generador = GeneradorCorreos({}, str(CATASTRO), despachador=lambda _: None)
    resultado = generador.procesar_espera(pd.DataFrame([{
        "RIT": "R-1", "TRIBUNAL": "LAJA", "NOMBRE": "Caso",
        "DERIVACION": _programa(catastro), "T ESPERA": float("inf"),
        "OBSERVACION": PATRON_ESPERA,
    }]))
    assert resultado["borradores_creados"] == 0
    assert any("T ESPERA" in error for error in resultado["errores"])


def test_alias_en_conflicto_no_sobrescribe_contacto(tmp_path):
    catastro = tmp_path / "catastro_programas.json"
    catastro.write_text(json.dumps([
        {"Nombre": "Programa Uno", "Mail": "uno@example.invalid", "Director": ""},
        {"Nombre": "Programa Dos", "Mail": "dos@example.invalid", "Director": ""},
    ]), encoding="utf-8")
    (tmp_path / "aliases_programas.json").write_text(json.dumps([{
        "alias": "Programa Dos", "nombre_catastro": "Programa Uno",
    }]), encoding="utf-8")

    contactos = CatastroContactos(str(catastro))
    assert contactos.resolver("Programa Uno")["mail"] == "uno@example.invalid"
    assert contactos.resolver("Programa Dos") is None


def test_contacto_duplicado_con_correos_distintos_es_ambiguo(tmp_path):
    catastro = tmp_path / "catastro_programas.json"
    catastro.write_text(json.dumps([
        {"Nombre": "Programa Uno", "Mail": "uno@example.invalid", "Director": ""},
        {"Nombre": "Programa Uno", "Mail": "otro@example.invalid", "Director": ""},
    ]), encoding="utf-8")
    contactos = CatastroContactos(str(catastro))
    assert contactos.resolver("Programa Uno") is None


def test_get_int_rechaza_decimales_e_infinitos():
    assert get_int("12") == 12
    assert get_int("12.0") == 12
    assert get_int("12.5") is None
    assert get_int(float("inf")) is None


def test_email_rechaza_inyeccion_de_destinatarios():
    assert GeneradorCorreos._mail_valido("programa@example.invalid")
    assert not GeneradorCorreos._mail_valido(
        "programa@example.invalid;atacante@example.invalid"
    )


def test_hoja2_duplicados_conserva_ultima_fila():
    """Catálogo C-10: 'Si hay duplicados en Hoja2 se conserva la última
    fila encontrada' — jamás se aborta el procesamiento por duplicados."""
    df = pd.DataFrame([
        {"rit": "R-1", "rut": "1-9", "nombre": "Caso", "tribunal": "LAJA",
         "programa": "PIE", "vence": datetime.now() + timedelta(days=10)},
        {"rit": "R-1", "rut": "1-9", "nombre": "Caso", "tribunal": "LAJA",
         "programa": "PIE", "vence": datetime.now() + timedelta(days=20)},
    ])
    cols = {campo: campo for campo in ("rit", "rut", "nombre", "tribunal", "programa")}
    cols["vencimiento"] = "vence"
    indice = _construir_indice_hoja2(df, cols)
    assert len(indice) == 1
    assert next(iter(indice.values())) == (datetime.now() + timedelta(days=20)).date()


def test_hoja2_malformada_degrada_a_cruce_desactivado():
    """Hoja2 sin columnas requeridas: cruce vacío con warning — el
    procesamiento de CUMPLIMIENTO no se aborta (Mejoras §2.2)."""
    df = pd.DataFrame([{"rit": "R-1"}])
    cols = {"rit": "rit"}  # faltan rut/nombre/tribunal/programa/vencimiento
    assert _construir_indice_hoja2(df, cols) == {}


def test_preview_aplica_validacion_bloqueante():
    df = pd.DataFrame([{
        "RIT": "R-1", "DERIVACION": "PIE", "TRIBUNAL": "DESCONOCIDO",
        "NOMBRE": "Caso", "FECHA VENCIMIENTO": datetime.now(),
    }])
    salida, error = calcular_preview(df, "INFORMES")
    assert salida is None
    assert "A8" in error


def test_exportacion_no_convierte_texto_en_formula(tmp_path):
    class Cola:
        def put(self, _mensaje):
            return None

    valores = [
        '=HYPERLINK("https://example.invalid")',
        "+1+1",
        "-1+1",
        "@SUM(A1:A2)",
        "  =1+1",
    ]
    nombre = _guardar_excel(
        pd.DataFrame({"NOMBRE": valores}),
        "PRUEBA",
        str(tmp_path),
        Cola(),
    )
    libro = load_workbook(tmp_path / nombre, data_only=False)
    for celda, original in zip(libro["PRUEBA"]["A2:A6"], valores, strict=True):
        celda = celda[0]
        # El texto se conserva EXACTO (sin prefijos visibles que corrompan
        # placeholders como '---') y nunca queda como fórmula ejecutable.
        assert celda.data_type != "f"
        assert celda.value == original


def test_informes_solo_agrega_audiencia_como_complementaria():
    hoy = datetime.now()
    cols = {
        "programa": "DERIVACION",
        "nombre": "NOMBRE",
        "vencimiento": "FECHA VENCIMIENTO",
        "curador": "CURADOR",
        "oido": "FEC. OIDO",
        "prox_aud": "PROXS. AUDS.",
        "rit": "RIT",
    }
    base = {
        "DERIVACION": "PIE Sintético",
        "NOMBRE": "Caso",
        "FECHA VENCIMIENTO": hoy - timedelta(days=1),
        "CURADOR": "NO POSEE",
        "FEC. OIDO": hoy,
        "PROXS. AUDS.": "",
        "RIT": "R-1",
    }
    sin_audiencia = generar_observacion_informes(
        pd.Series(base), "LAJA", cols, fila_excel=2
    )
    assert "curador" not in sin_audiencia.lower()
    assert "oído" not in sin_audiencia.lower()

    base["PROXS. AUDS."] = hoy + timedelta(days=1)
    con_audiencia = generar_observacion_informes(
        pd.Series(base), "LAJA", cols, fila_excel=2
    )
    assert "Se cita a audiencia" in con_audiencia


def test_validador_de_archivo_rechaza_tipo_tamano_y_xlsx_falso(tmp_path):
    texto = tmp_path / "entrada.txt"
    texto.write_text("dato", encoding="utf-8")
    with pytest.raises(ValueError, match="Extensión"):
        validar_archivo_excel(texto)

    grande = tmp_path / "grande.xls"
    grande.write_bytes(b"12345")
    with pytest.raises(ValueError, match="tamaño"):
        validar_archivo_excel(grande, max_bytes=4)

    falso = tmp_path / "falso.xlsx"
    falso.write_text("no es zip", encoding="utf-8")
    with pytest.raises(ValueError, match="XLSX/XLSM"):
        validar_archivo_excel(falso)


def test_validar_archivo_excel_rechaza_dimensiones_excesivas(tmp_path):
    from openpyxl import Workbook

    from motor.utilidades import validar_archivo_excel

    ruta = tmp_path / "grande.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(row=4, column=1, value="ok")
    wb.save(ruta)

    assert validar_archivo_excel(ruta, max_filas=10, max_columnas=10) == ruta

    ws.cell(row=11, column=1, value="demasiadas filas")
    wb.save(ruta)
    with pytest.raises(ValueError, match="supera el máximo permitido"):
        validar_archivo_excel(ruta, max_filas=10, max_columnas=10)


def test_exportar_borrador_html_escapa_metadatos_y_es_atomico(tmp_path):
    from comunicaciones.generador_correos import exportar_borrador_html

    assert exportar_borrador_html({
        "para": ["programa@example.invalid"],
        "cc": ["ucc@example.invalid"],
        "asunto": "<Asunto crítico>",
        "cuerpo_html": "<p>Cuerpo autorizado</p>",
    }, str(tmp_path))

    archivos = list(tmp_path.glob("*.html"))
    assert len(archivos) == 1
    contenido = archivos[0].read_text(encoding="utf-8")
    assert "&lt;Asunto crítico&gt;" in contenido
    assert "<p>Cuerpo autorizado</p>" in contenido
    assert not list(tmp_path.glob("*.tmp.html"))
